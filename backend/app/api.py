from __future__ import annotations

import asyncio
import json
from pathlib import Path
from typing import Optional

from fastapi import FastAPI, File, Query, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse

from app.schemas import IndexResponse, UploadResponse
from app.services.chat_service import (
    generate_session_title,
    stream_answer_tokens_with_history,
)
from app.services.file_service import (
    ensure_app_dirs,
    get_doc_path,
    list_docs,
    save_uploaded_doc,
)
from app.services.queue_service import CHAT_QUEUE
from app.services.session_store import (
    append_message,
    list_sessions,
    load_session,
    new_session_id,
    save_session,
)
from app.services.state import (
    cleanup_finished_index_jobs,
    create_index_job,
    fail_index_job,
    finish_index_job,
    get_embedder,
    get_index_job,
    get_llm,
    get_store,
    reset_store_cache,
    start_index_job,
    update_index_job,
)

from app.services.chat_service import (
    generate_session_title,
    stream_answer_tokens_with_history,
    generate_example_prompts,
)

from src.config import CONFIG
from src.ingest import ingest_docs, reindex_single_doc_keep_others


app = FastAPI(title="ChatDoc API")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:3000"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


def sse(event: str, data: dict) -> str:
    return f"event: {event}\ndata: {json.dumps(data, ensure_ascii=False)}\n\n"


def toast(message: str, kind: str = "info") -> str:
    return sse("toast", {"message": message, "kind": kind})


@app.on_event("startup")
def _startup() -> None:
    try:
        ensure_app_dirs()
    except Exception:
        pass

    for fn in (get_store, get_embedder, get_llm):
        try:
            fn()
        except Exception:
            pass


@app.get("/health")
def health():
    return {"ok": True}

@app.get("/api/example-prompts")
def api_example_prompts(
    selected_pdf: str = Query("Semua dokumen", description="Semua dokumen atau nama file dokumen"),
):
    ensure_app_dirs()
    examples = generate_example_prompts(selected_doc=selected_pdf, max_examples=4)
    return {"examples": examples}



def _indexed_doc_names() -> set[str]:
    try:
        p = Path(CONFIG.DOC_REGISTRY_PATH)
        if not p.exists():
            return set()

        data = json.loads(p.read_text(encoding="utf-8"))
        docs = (data or {}).get("docs", {}) or {}
        return {v.get("doc_name") for v in docs.values() if (v or {}).get("doc_name")}
    except Exception:
        return set()


def _wipe_index_files() -> None:
    for p in (Path(CONFIG.CHUNKS_PATH), Path(CONFIG.DOC_REGISTRY_PATH)):
        try:
            if p.exists():
                p.unlink()
        except Exception:
            pass

    try:
        reset_store_cache()
    except Exception:
        pass


def _doc_path_by_name(name: str) -> Optional[Path]:
    for p in list_docs():
        if p.name == name:
            return p
    return None


def _ensure_index_for_docs(target_docs: list[str], force: bool = False) -> dict:
    ensure_app_dirs()

    docs = list_docs()
    if not docs:
        return {
            "status": "empty",
            "added_docs": 0,
            "added_chunks": 0,
            "message": "Tidak ada dokumen untuk di-index.",
        }

    doc_map = {p.name: p for p in docs}
    wanted_paths = [doc_map[n] for n in target_docs if n in doc_map]
    if not wanted_paths:
        return {
            "status": "empty",
            "added_docs": 0,
            "added_chunks": 0,
            "message": "Target dokumen tidak ditemukan.",
        }

    if force:
        _wipe_index_files()
        result = ingest_docs(wanted_paths)
        try:
            reset_store_cache()
        except Exception:
            pass
        return result

    indexed = _indexed_doc_names()
    missing_paths = [doc_map[n] for n in target_docs if (n in doc_map and n not in indexed)]
    if not missing_paths:
        return {"status": "ok", "added_docs": 0, "added_chunks": 0}

    result = ingest_docs(missing_paths)
    try:
        reset_store_cache()
    except Exception:
        pass
    return result


def _resolve_index_targets(scope: str, force: bool) -> tuple[list[Path], Optional[str], bool]:
    """
    Return:
    - target_paths
    - error_message
    - single_reindex_mode
    """
    docs = list_docs()
    all_names = [p.name for p in docs]
    doc_map = {p.name: p for p in docs}

    if not all_names:
        return [], "Tidak ada dokumen untuk di-index.", False

    if scope == "Semua dokumen":
        if force:
            return [doc_map[n] for n in all_names], None, False

        indexed = _indexed_doc_names()
        missing_paths = [doc_map[n] for n in all_names if n not in indexed]
        return missing_paths, None, False

    if scope not in all_names:
        return [], "Dokumen scope tidak ditemukan.", False

    if force:
        return [doc_map[scope]], None, True

    indexed = _indexed_doc_names()
    if scope in indexed:
        return [], "Dokumen sudah ter-index.", False

    return [doc_map[scope]], None, False


@app.get("/api/docs")
def api_list_docs():
    ensure_app_dirs()
    out = []
    for p in list_docs():
        ext = p.suffix.lower()
        out.append(
            {
                "name": p.name,
                "type": ext.lstrip("."),
                "size": p.stat().st_size if p.exists() else None,
                "modified": int(p.stat().st_mtime) if p.exists() else None,
            }
        )
    return {"docs": out}


@app.get("/api/pdfs")
def api_list_pdfs():
    ensure_app_dirs()
    return {"pdfs": [p.name for p in list_docs()]}


@app.post("/api/pdf/upload", response_model=UploadResponse)
async def upload_pdf(file: UploadFile = File(...)):
    ensure_app_dirs()
    content = await file.read()
    out_path = save_uploaded_doc(file.filename, content)
    return UploadResponse(doc_name=out_path.name, doc_path=str(out_path))


@app.get("/api/pdf/{doc_name}")
def serve_pdf(doc_name: str):
    p = get_doc_path(doc_name)
    ext = p.suffix.lower()
    media_type = {
        ".pdf": "application/pdf",
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    }.get(ext, "application/octet-stream")
    return FileResponse(str(p), media_type=media_type, filename=p.name)


@app.get("/api/doc/{doc_name}/preview")
def preview_doc(doc_name: str):
    """
    Preview endpoint untuk frontend.
    - PDF: tidak di-preview lewat endpoint ini (gunakan /api/pdf/{doc_name})
    - DOCX: return text blocks sederhana
    """
    ensure_app_dirs()
    p = get_doc_path(doc_name)
    ext = p.suffix.lower()

    if ext == ".pdf":
        return JSONResponse(
            {"type": "pdf", "message": "Gunakan /api/pdf/{doc_name} untuk PDF viewer."}
        )

    if ext == ".docx":
        from docx import Document

        d = Document(str(p))
        paras = []
        for i, para in enumerate(d.paragraphs, start=1):
            t = (para.text or "").strip()
            if t:
                paras.append({"i": i, "text": t})

        tables = []
        for ti, table in enumerate(d.tables, start=1):
            rows = []
            for r in table.rows:
                rows.append([(c.text or "").strip() for c in r.cells])
            tables.append({"i": ti, "rows": rows})

        return JSONResponse(
            {
                "type": "docx",
                "name": p.name,
                "paragraphs": paras[:800],
                "tables": tables[:50],
            }
        )

    return JSONResponse({"type": "unknown", "name": p.name})


@app.get("/api/index/status")
def index_status():
    ensure_app_dirs()
    doc_names = [p.name for p in list_docs()]
    indexed = sorted(_indexed_doc_names())
    missing = sorted([n for n in doc_names if n not in set(indexed)])
    return {"docs": doc_names, "indexed": indexed, "missing": missing}


@app.post("/api/index", response_model=IndexResponse)
def build_index(
    force: bool = Query(False, description="True untuk re-index"),
    scope: str = Query("Semua dokumen", description="Semua dokumen atau nama file dokumen"),
):
    ensure_app_dirs()

    docs = list_docs()
    all_names = [p.name for p in docs]
    if not all_names:
        return IndexResponse(
            status="empty",
            added_docs=0,
            added_chunks=0,
            message="Tidak ada dokumen untuk di-index.",
        )

    if scope == "Semua dokumen":
        result = _ensure_index_for_docs(all_names, force=force)
        return IndexResponse(
            status=result.get("status", "ok"),
            added_docs=int(result.get("added_docs", 0)),
            added_chunks=int(result.get("added_chunks", 0)),
            message=result.get("message"),
        )

    if scope not in all_names:
        return IndexResponse(
            status="empty",
            added_docs=0,
            added_chunks=0,
            message="Dokumen scope tidak ditemukan.",
        )

    if not force:
        indexed = _indexed_doc_names()
        if scope in indexed:
            return IndexResponse(status="ok", added_docs=0, added_chunks=0, message="Dokumen sudah ter-index.")

        result = _ensure_index_for_docs([scope], force=False)
        return IndexResponse(
            status=result.get("status", "ok"),
            added_docs=int(result.get("added_docs", 0)),
            added_chunks=int(result.get("added_chunks", 0)),
            message=result.get("message"),
        )

    p = _doc_path_by_name(scope)
    if p is None:
        return IndexResponse(
            status="empty",
            added_docs=0,
            added_chunks=0,
            message="Dokumen tidak ditemukan.",
        )

    result = reindex_single_doc_keep_others(p)

    try:
        reset_store_cache()
    except Exception:
        pass

    return IndexResponse(
        status=result.get("status", "ok"),
        added_docs=int(result.get("added_docs", 0)),
        added_chunks=int(result.get("added_chunks", 0)),
        message=result.get("message"),
    )


@app.get("/api/index/stream")
async def index_stream(
    force: bool = Query(False, description="True untuk re-index"),
    scope: str = Query("Semua dokumen", description="Semua dokumen atau nama file dokumen"),
):
    async def gen():
        ensure_app_dirs()
        cleanup_finished_index_jobs()

        target_paths, error_message, single_reindex_mode = _resolve_index_targets(scope, force)
        total_docs = len(target_paths) if target_paths else (1 if single_reindex_mode else 0)

        job = create_index_job(
            scope=scope,
            force=force,
            total_docs=total_docs,
            title="Re-indexing documents" if force else "Indexing documents",
        )
        job_id = job["job_id"]

        yield sse("start", job)
        yield toast("Memulai proses indexing…", "info")

        if error_message:
            fail_payload = fail_index_job(job_id, error_message) or {"job_id": job_id, "message": error_message}
            kind = "warning" if error_message in {"Dokumen sudah ter-index.", "Tidak ada dokumen untuk di-index."} else "error"
            yield toast(error_message, kind)
            yield sse("error", fail_payload)
            return

        if not target_paths:
            done_payload = finish_index_job(
                job_id,
                message="Tidak ada dokumen baru untuk di-index.",
                added_docs=0,
                added_chunks=0,
            ) or {"job_id": job_id, "message": "Tidak ada dokumen baru untuk di-index."}
            yield toast("Tidak ada dokumen baru untuk di-index.", "success")
            yield sse("done", done_payload)
            return

        start_index_job(
            job_id,
            title="Re-indexing documents" if force else "Indexing documents",
            message="Preparing indexing pipeline",
        )

        current_loop = asyncio.get_running_loop()
        progress_queue: asyncio.Queue[dict] = asyncio.Queue()

        def progress_cb(**payload):
            updated = update_index_job(job_id, **payload)
            if updated is None:
                updated = get_index_job(job_id) or {"job_id": job_id, **payload}
            current_loop.call_soon_threadsafe(progress_queue.put_nowait, updated)

        def run_index():
            try:
                if force and single_reindex_mode:
                    result = reindex_single_doc_keep_others(
                        target_paths[0],
                        progress_cb=progress_cb,
                    )
                else:
                    if force and scope == "Semua dokumen":
                        _wipe_index_files()
                    result = ingest_docs(
                        target_paths,
                        progress_cb=progress_cb,
                    )

                try:
                    reset_store_cache()
                except Exception:
                    pass

                final_message = result.get("message")
                if not final_message:
                    final_message = (
                        "Re-index completed successfully."
                        if force
                        else "Indexing completed successfully."
                    )

                finished = finish_index_job(
                    job_id,
                    message=final_message,
                    added_docs=int(result.get("added_docs", 0)),
                    added_chunks=int(result.get("added_chunks", 0)),
                )
                current_loop.call_soon_threadsafe(
                    progress_queue.put_nowait,
                    finished or {"job_id": job_id, "status": "done", "message": final_message},
                )
            except Exception as e:
                failed = fail_index_job(job_id, str(e))
                current_loop.call_soon_threadsafe(
                    progress_queue.put_nowait,
                    failed or {"job_id": job_id, "status": "error", "message": str(e)},
                )

        task = current_loop.run_in_executor(None, run_index)

        last_progress = -1
        last_stage = None
        done_sent = False

        while True:
            if task.done() and progress_queue.empty():
                final_job = get_index_job(job_id)
                if final_job:
                    status = final_job.get("status")
                    if status == "done" and not done_sent:
                        yield toast(final_job.get("message") or "Indexing selesai.", "success")
                        yield sse("done", final_job)
                        done_sent = True
                    elif status == "error":
                        yield toast(final_job.get("message") or "Indexing gagal.", "error")
                        yield sse("error", final_job)
                break

            try:
                payload = await asyncio.wait_for(progress_queue.get(), timeout=0.35)
            except asyncio.TimeoutError:
                heartbeat = get_index_job(job_id)
                if heartbeat and heartbeat.get("status") == "running":
                    yield sse(
                        "progress",
                        {
                            "job_id": heartbeat.get("job_id"),
                            "status": heartbeat.get("status"),
                            "stage": heartbeat.get("stage"),
                            "progress": heartbeat.get("progress"),
                            "message": heartbeat.get("message"),
                            "current_document": heartbeat.get("current_document"),
                            "current_page": heartbeat.get("current_page"),
                            "total_pages": heartbeat.get("total_pages"),
                            "processed_docs": heartbeat.get("processed_docs"),
                            "total_docs": heartbeat.get("total_docs"),
                            "processed_chunks": heartbeat.get("processed_chunks"),
                            "total_chunks": heartbeat.get("total_chunks"),
                            "processed_batches": heartbeat.get("processed_batches"),
                            "total_batches": heartbeat.get("total_batches"),
                        },
                    )
                continue

            status = payload.get("status")
            stage = payload.get("stage")
            progress = int(payload.get("progress", 0) or 0)

            if status == "running":
                yield sse("progress", payload)

                if stage != last_stage or progress != last_progress:
                    if stage == "chunking":
                        yield toast(payload.get("message") or "Sedang mengekstrak dokumen…", "info")
                    elif stage == "embedding":
                        yield toast(payload.get("message") or "Sedang membuat embedding…", "info")
                    elif stage == "storing":
                        yield toast(payload.get("message") or "Sedang menyimpan ke vector store…", "info")
                    elif stage == "finalizing":
                        yield toast(payload.get("message") or "Menyelesaikan indexing…", "info")

                last_stage = stage
                last_progress = progress

            elif status == "done":
                if not done_sent:
                    yield toast(payload.get("message") or "Indexing selesai.", "success")
                    yield sse("done", payload)
                    done_sent = True
                break

            elif status == "error":
                yield toast(payload.get("message") or "Indexing gagal.", "error")
                yield sse("error", payload)
                break

    return StreamingResponse(
        gen(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no",
        },
    )


@app.post("/api/session/new")
def api_new_session():
    sid = new_session_id()
    sess = {"session_id": sid, "title": "Obrolan Baru", "messages": []}
    save_session(sess)
    return {"session_id": sid}


@app.get("/api/sessions")
def api_sessions(limit: int = 50):
    return {"sessions": list_sessions(limit=limit)}


@app.get("/api/session/{session_id}")
def api_get_session(session_id: str):
    return load_session(session_id)


@app.get("/api/chat/stream")
async def chat_stream(
    message: str,
    selected_pdf: str = "Semua dokumen",
    session_id: str = Query(..., description="Chat session id"),
):
    selected_doc = selected_pdf or "Semua dokumen"

    async def gen():
        sess = load_session(session_id)
        if "session_id" not in sess:
            yield toast("Session tidak valid.", "error")
            yield sse("error", {"message": "Invalid session"})
            return

        try:
            ensure_app_dirs()
            all_names = [p.name for p in list_docs()]
            indexed = _indexed_doc_names()

            if selected_doc == "Semua dokumen":
                targets = all_names
            else:
                targets = [selected_doc] if selected_doc in all_names else all_names

            missing = [n for n in targets if n not in indexed]
            if missing:
                yield toast(
                    "Dokumen belum di-index. Silakan lakukan indexing terlebih dahulu.",
                    "warning",
                )
                yield sse(
                    "index_required",
                    {"message": "Dokumen belum di-index.", "missing": missing},
                )
                return
        except Exception as e:
            yield toast(f"Gagal cek index: {str(e)}", "error")
            yield sse("error", {"message": f"Gagal cek index: {str(e)}"})
            return

        try:
            sess0 = load_session(session_id)
            msgs0 = sess0.get("messages", []) if isinstance(sess0, dict) else []
            has_user_msg = any((m.get("role") == "user") for m in msgs0)

            if not has_user_msg:
                title = generate_session_title(message)
                sess0["title"] = title
                save_session(sess0)
                yield sse("session_title", {"session_id": session_id, "title": title})
        except Exception:
            pass

        append_message(session_id, "user", message)

        sess = load_session(session_id)
        history = sess.get("messages", [])

        job_id = CHAT_QUEUE.enqueue(message=message, selected_pdf=selected_doc)

        last_pos = None
        while True:
            pos = CHAT_QUEUE.get_position(job_id)
            if pos is None:
                yield toast("Job tidak ditemukan di antrian.", "error")
                yield sse("error", {"message": "Job not found in queue"})
                return

            if pos == 0:
                yield toast("Mulai memproses…", "info")
                yield sse("start", {"job_id": job_id, "session_id": session_id})
                break

            if pos != last_pos:
                yield toast(f"Masuk antrian: posisi #{pos}", "info")
                last_pos = pos

            yield sse("queued", {"job_id": job_id, "position": pos, "session_id": session_id})
            await asyncio.sleep(0.5)

        loop = asyncio.get_running_loop()
        job = await loop.run_in_executor(None, CHAT_QUEUE.pop_next)

        if job.job_id != job_id:
            yield toast("Queue race condition, silakan ulangi.", "error")
            yield sse("error", {"message": "Queue race condition, please retry"})
            return

        answer_buf: list[str] = []
        llm_eval: dict = {}
        try:
            stream_result = stream_answer_tokens_with_history(
                message=job.message,
                history=history,
                selected_doc=job.selected_pdf,
            )

            if isinstance(stream_result, tuple) and len(stream_result) >= 3:
                token_iter, sources, llm_eval = stream_result[0], stream_result[1], stream_result[2]
            else:
                token_iter, sources = stream_result
                llm_eval = {}

            for token in token_iter:
                answer_buf.append(token)
                yield sse("token", {"token": token})
                await asyncio.sleep(0)

            full_answer = "".join(answer_buf).strip()

            if isinstance(llm_eval, dict):
                llm_eval["output_chars"] = len(full_answer)
                if "output_tokens_est" not in llm_eval:
                    llm_eval["output_tokens_est"] = max(1, len(full_answer) // 4) if full_answer else 0

            append_message(session_id, "assistant", full_answer, citations=sources)

            yield toast("Selesai.", "success")

            if llm_eval:
                yield sse(
                    "llm_eval",
                    {
                        "job_id": job_id,
                        "session_id": session_id,
                        "eval": llm_eval,
                    },
                )

            yield sse(
                "done",
                {
                    "job_id": job_id,
                    "session_id": session_id,
                    "citations": sources,
                    "llm_eval": llm_eval,
                },
            )

        except asyncio.CancelledError:
            return
        except Exception as e:
            yield toast(str(e), "error")
            yield sse(
                "error",
                {"job_id": job_id, "session_id": session_id, "message": str(e)},
            )
        finally:
            try:
                CHAT_QUEUE.mark_done()
            except Exception:
                pass

    return StreamingResponse(
        gen(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no",
        },
    )

